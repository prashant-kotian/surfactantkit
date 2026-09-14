"""Renders round3_questions.json into a human-readable Excel answer key."""
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment

HERE = Path(__file__).resolve().parent
questions = json.loads((HERE / "round3_questions.json").read_text(encoding="utf-8"))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Round3 Answer Key"

headers = ["ID", "SMILES", "Data provenance", "charge_type", "cmc_mM", "gamma_max_mol_per_m2",
           "a_min_nm2", "isotherm_model", "frumkin_a", "delta_g_mic_kJ_per_mol", "gaps (toolkit's own)"]
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

for q in questions:
    g = q["gold"]
    ws.append([
        q["id"], q["smiles"], q["data_provenance"], g["charge_type"], g["cmc_mM"], g["gamma_max_mol_per_m2"],
        g["a_min_nm2"], g["isotherm_model"], g["frumkin_a"], g["delta_g_mic_kJ_per_mol"],
        "\n".join(g["gaps"]),
    ])

for col_idx, width in zip(range(1, 12), [8, 40, 60, 14, 12, 16, 10, 12, 10, 16, 80]):
    ws.column_dimensions[chr(64 + col_idx)].width = width
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(HERE / "Round3_Answer_Key.xlsx")
print("wrote Round3_Answer_Key.xlsx")
